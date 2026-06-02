import streamlit as st
import pandas as pd
import altair as alt
import os

import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows

st.set_page_config(page_title="Two-Phase Timeline Creator", layout="wide")
st.title("Video Timeline Creator")

# 1. File Configuration
SHEET_NAME = 'Subset_TimeStamps'

VIDEO_FILE = st.session_state.get('selected_video', None)
EXCEL_FILE = st.session_state.get('selected_excel', None)

if not VIDEO_FILE or not os.path.exists(VIDEO_FILE):
    st.error("No valid video selected. Please return to the Home page.")
    st.stop()

if not EXCEL_FILE or not os.path.exists(EXCEL_FILE):
    st.error("No valid Excel file selected. Please return to the Home page.")
    st.stop()

# 2. Load Data
@st.cache_data
def load_data(filepath):
    if os.path.exists(filepath):
        try:
            xls = pd.ExcelFile(filepath)
            actual_sheets = xls.sheet_names
            
            target_sheet = None
            for sheet in actual_sheets:
                if sheet.strip() == SHEET_NAME:
                    target_sheet = sheet
                    break
            
            if target_sheet:
                return pd.read_excel(filepath, sheet_name=target_sheet)
            else:
                st.error(f"Target sheet '{SHEET_NAME}' not found.")
                st.info(f"Sheets available in **{filepath}**: {', '.join(actual_sheets)}")
                st.stop()
                
        except Exception as e:
            st.error(f"Error reading Excel file: {e}")
            st.stop()
            
    return pd.DataFrame()

if 'df' not in st.session_state:
    st.session_state.df = load_data(EXCEL_FILE)

df = st.session_state.df

# Standardize all column headers to uppercase to prevent case-sensitivity KeyError
df.columns = df.columns.str.upper()

# Check for the required two-phase columns
has_two_phases = all(col in df.columns for col in ['PRE_SNAP', 'POST_SNAP'])

if not df.empty and has_two_phases:
    
    st.caption(f"Analyzing: **{VIDEO_FILE}** with data from **{EXCEL_FILE}**")
    
    valid_plays = df.copy() 
    
    # Calculate the next play's pre-snap time to define when the "Post-Snap" phase ends
    valid_plays['NEXT_PRE_SNAP'] = valid_plays['PRE_SNAP'].shift(-1)
    
    if 'active_play_id' not in st.session_state:
        st.session_state.active_play_id = valid_plays['PLAY_UNIQUE_ID'].iloc[0]
    if 'active_phase' not in st.session_state:
        st.session_state.active_phase = "Pre-Snap"
    
    col_video, col_controls = st.columns([2.5, 1])
    
    with col_controls:
        # Reserve a spot at the very top of the column for the scoreboard
        scoreboard_container = st.container()
        
        st.subheader("Play Controls")
        
        play_options = valid_plays['PLAY_UNIQUE_ID'].tolist()
        current_play_index = play_options.index(st.session_state.active_play_id) if st.session_state.active_play_id in play_options else 0
        
        selected_play_id = st.selectbox("Select Play ID", play_options, index=current_play_index)
        
        phase_options = ["Pre-Snap", "Post-Snap"]
        current_phase_index = phase_options.index(st.session_state.active_phase)
        selected_phase = st.selectbox("Select Phase", phase_options, index=current_phase_index)

        if selected_play_id != st.session_state.active_play_id or selected_phase != st.session_state.active_phase:
            st.session_state.active_play_id = selected_play_id
            st.session_state.active_phase = selected_phase
            st.rerun()

        row = valid_plays[valid_plays['PLAY_UNIQUE_ID'] == st.session_state.active_play_id].iloc[0]

        # Display the event name
        st.markdown(f"**Event:** {row.get('EVENT_NAME', 'N/A')}")
        
        # Safely extract timestamps using .get()
        raw_pre = row.get('PRE_SNAP', 0)
        raw_post = row.get('POST_SNAP', 0)
        raw_next_pre = row.get('NEXT_PRE_SNAP', pd.NA)
        
        # Determine jump points based on the active phase starting at the phase time
        if st.session_state.active_phase == 'Pre-Snap':
            start_val = raw_pre if pd.notna(raw_pre) else 0
            end_val = raw_post
        else: # Post-Snap
            start_val = raw_post if pd.notna(raw_post) else 0
            # If it's the last play in the sheet, add a 10 second buffer
            end_val = raw_next_pre if pd.notna(raw_next_pre) else start_val + 10 

        start_time = int(start_val) if pd.notna(start_val) else 0
        end_time = int(end_val) if pd.notna(end_val) else 0

        # --- Scoreboard Banner ---
        # Safely extract scores
        raw_home_score = row.get('HOME_SCORE', 0)
        home_score = int(float(raw_home_score)) if pd.notna(raw_home_score) and str(raw_home_score).strip() != "" else 0
        
        raw_away_score = row.get('AWAY_SCORE', 0)
        away_score = int(float(raw_away_score)) if pd.notna(raw_away_score) and str(raw_away_score).strip() != "" else 0

        home_team = row.get('HOME_TEAM_NAME', 'CHI')
        away_team = row.get('AWAY_TEAM_NAME', 'MIN')
        
        # Safely extract Down and Distance to prevent "nan"
        raw_down = row.get('DOWN', 'N/A')
        clean_down = str(raw_down).replace(".0", "") if pd.notna(raw_down) and str(raw_down).lower() != "nan" else "N/A"
        
        # Add the proper suffix to the down
        ordinal_map = {"1": "1st", "2": "2nd", "3": "3rd", "4": "4th"}
        down = ordinal_map.get(clean_down, clean_down) # Defaults to the plain value if it's "N/A" or unusual
        
        raw_dist = row.get('DISTANCE', 'N/A')
        distance = str(raw_dist).replace(".0", "") if pd.notna(raw_dist) and str(raw_dist).lower() != "nan" else "N/A"

        scoreboard_html = f"""
        <div style="display: flex; justify-content: space-around; align-items: center; 
                    background-color: #1E1E2E; color: white; padding: 15px; 
                    border-radius: 10px; border: 1px solid #333; margin-bottom: 20px;">
            
            <div style="text-align: center;">
                <span style="font-size: 12px; color: #888;">AWAY</span><br>
                <span style="font-size: 24px; font-weight: bold;">{away_team} {away_score}</span>
            </div>
            
            <div style="font-size: 18px; color: #888;">@</div>
            
            <div style="text-align: center;">
                <span style="font-size: 12px; color: #888;">HOME</span><br>
                <span style="font-size: 24px; font-weight: bold;">{home_team} {home_score}</span>
            </div>
            
            <div style="text-align: center; background-color: #2A2A3C; padding: 10px 15px; border-radius: 5px;">
                <span style="font-size: 16px; font-weight: bold; color: #FFA500;">{down} & {distance}</span>
            </div>
        </div>
        """
        # Inject the HTML into the container we reserved at the top
        scoreboard_container.html(scoreboard_html)
        # -----------------------
        
        if start_time > 0 or end_time > 0:
            st.markdown(f"*Playhead will jump to: {start_time}s*")
        else:
            st.markdown("*Playhead will start at 0s (Pending Timestamps)*")
            
        with st.expander("Expected Metrics"):
            
            # Helper function to format metrics safely
            def format_metric(val, is_percentage=False):
                if pd.isna(val) or str(val).strip().lower() == "nan" or str(val).strip() == "":
                    return "N/A"
                try:
                    num = float(val)
                    if is_percentage:
                        return f"{num * 100:.4f}%"
                    # For non-percentages (like yards), just round to 2 decimal places
                    return f"{num:.2f}"
                except (ValueError, TypeError):
                    return str(val)

            col_m1, col_m2 = st.columns(2)
            with col_m1:
                st.markdown(f"**xYards:** {format_metric(row.get('XYARDS', 'N/A'))}")
                st.markdown(f"**xRush%:** {format_metric(row.get('XRUSH%', 'N/A'), is_percentage=True)}")
                st.markdown(f"**xOvrYds:** {format_metric(row.get('XOVRYDS', 'N/A'))}")
                st.markdown(f"**xSack%:** {format_metric(row.get('XSACK%', 'N/A'), is_percentage=True)}")
            with col_m2:
                st.markdown(f"**xScramble%:** {format_metric(row.get('XSCRAMBLE%', 'N/A'), is_percentage=True)}")
                st.markdown(f"**xTD%:** {format_metric(row.get('XTD%', 'N/A'), is_percentage=True)}")
                st.markdown(f"**xRushTD%:** {format_metric(row.get('XRUSHTD%', 'N/A'), is_percentage=True)}")

    with col_video:
        st.subheader("Full Video Playback")
        st.video(VIDEO_FILE, start_time=start_time)
        st.caption(f"Playing full video file. Playhead moved to {start_time}s.")

    st.divider()

    st.subheader("Scheme & Personnel")
    
    # Safely extract data, handling blank cells
    raw_off_form = row.get('OFFENSIVE_FORMATION', 'N/A')
    off_form = str(raw_off_form) if pd.notna(raw_off_form) and str(raw_off_form).lower() != "nan" else "N/A"
    
    raw_def_front = row.get('DEFENSIVE_FRONT', 'N/A')
    def_front = str(raw_def_front) if pd.notna(raw_def_front) and str(raw_def_front).lower() != "nan" else "N/A"
    
    raw_off_pers = row.get('OFFENSIVE_PERSONNEL', 'N/A')
    off_pers = str(raw_off_pers).replace(".0", "") if pd.notna(raw_off_pers) and str(raw_off_pers).lower() != "nan" else "N/A"
    
    raw_def_pers = row.get('DEFENSIVE_PERSONNEL', 'N/A')
    def_pers = str(raw_def_pers) if pd.notna(raw_def_pers) and str(raw_def_pers).lower() != "nan" else "N/A"

    scheme_html = f"""
    <div style="display: flex; justify-content: space-between; gap: 15px; margin-bottom: 10px;">
        <div style="flex: 1; text-align: center; background-color: #1E1E2E; padding: 12px; border-radius: 8px; border: 1px solid #333;">
            <div style="font-size: 11px; color: #888; text-transform: uppercase; margin-bottom: 4px;">Off Formation</div>
            <div style="font-size: 16px; font-weight: bold; color: #FFA500;">{off_form}</div>
        </div>
        <div style="flex: 1; text-align: center; background-color: #1E1E2E; padding: 12px; border-radius: 8px; border: 1px solid #333;">
            <div style="font-size: 11px; color: #888; text-transform: uppercase; margin-bottom: 4px;">Def Front</div>
            <div style="font-size: 16px; font-weight: bold; color: #BA55D3;">{def_front}</div>
        </div>
        <div style="flex: 1; text-align: center; background-color: #1E1E2E; padding: 12px; border-radius: 8px; border: 1px solid #333;">
            <div style="font-size: 11px; color: #888; text-transform: uppercase; margin-bottom: 4px;">Off Personnel</div>
            <div style="font-size: 16px; font-weight: bold; color: #FFA500;">{off_pers}</div>
        </div>
        <div style="flex: 1; text-align: center; background-color: #1E1E2E; padding: 12px; border-radius: 8px; border: 1px solid #333;">
            <div style="font-size: 11px; color: #888; text-transform: uppercase; margin-bottom: 4px;">Def Personnel</div>
            <div style="font-size: 16px; font-weight: bold; color: #BA55D3;">{def_pers}</div>
        </div>
    </div>
    """
    st.html(scheme_html)

    # --- Full Roster Personnel Section ---
    
    # Determine which team is on Offense vs Defense for the headers
    raw_poss = row.get('POSSESSION_TEAM', away_team)
    poss_team = str(raw_poss) if pd.notna(raw_poss) and str(raw_poss).lower() != 'nan' else away_team
    def_team = home_team if poss_team == away_team else away_team

    def get_players(side):
        players = []
        for i in range(1, 12):
            fname = row.get(f"{side}_{i}_PLAYER_FIRST_NAME", pd.NA)
            lname = row.get(f"{side}_{i}_PLAYER_LAST_NAME", pd.NA)
            role = row.get(f"{side}_{i}_ROLE", "N/A")
            
            if pd.notna(fname) and str(fname).strip() not in ["nan", ""]:
                safe_lname = str(lname).strip() if pd.notna(lname) and str(lname).strip() != "nan" else ""
                full_name = f"{str(fname).strip()} {safe_lname}".strip()
                players.append((role, full_name))
        return players

    off_players = get_players('OFF')
    def_players = get_players('DEF')

    # Helper function to generate HTML rows for the players
    def build_player_rows(player_list):
        html = ""
        for role, name in player_list:
            html += f"<div style='margin-bottom: 6px;'><span style='color: #888; font-size: 11px; font-weight: bold; width: 35px; display: inline-block;'>{role}</span> <span style='font-size: 12px; color: #E0E0E0;'>{name}</span></div>"
        return html

    # Split the lists in half (6 and 5) to create the 2 columns per team
    off_col1 = build_player_rows(off_players[:6])
    off_col2 = build_player_rows(off_players[6:])
    
    def_col1 = build_player_rows(def_players[:6])
    def_col2 = build_player_rows(def_players[6:])

    roster_html = f"""
    <div style="display: flex; justify-content: space-between; gap: 15px; margin-bottom: 20px;">
        <div style="flex: 1; background-color: #1E1E2E; padding: 15px; border-radius: 8px; border: 1px solid #333;">
            <div style="font-size: 12px; font-weight: bold; color: #FFA500; text-transform: uppercase; margin-bottom: 10px; border-bottom: 1px solid #333; padding-bottom: 5px;">
                {poss_team} Personnel
            </div>
            <div style="display: flex; gap: 10px;">
                <div style="flex: 1;">{off_col1}</div>
                <div style="flex: 1;">{off_col2}</div>
            </div>
        </div>
        
        <div style="flex: 1; background-color: #1E1E2E; padding: 15px; border-radius: 8px; border: 1px solid #333;">
            <div style="font-size: 12px; font-weight: bold; color: #BA55D3; text-transform: uppercase; margin-bottom: 10px; border-bottom: 1px solid #333; padding-bottom: 5px;">
                {def_team} Personnel
            </div>
            <div style="display: flex; gap: 10px;">
                <div style="flex: 1;">{def_col1}</div>
                <div style="flex: 1;">{def_col2}</div>
            </div>
        </div>
    </div>
    """
    st.html(roster_html)

    st.subheader("Edit Timestamps")
    
    def highlight_active_row(row):
        styles = [''] * len(row)
        if float(row['PLAY_UNIQUE_ID']) == float(st.session_state.active_play_id):
            
            # Map the selected phase to the target column
            if st.session_state.active_phase == 'Pre-Snap':
                target_col = 'PRE_SNAP'
            else:
                target_col = 'POST_SNAP'
                
            for i, col_name in enumerate(row.index):
                if col_name == target_col:
                    styles[i] = 'background-color: #ff4b4b; font-weight: bold; color: white;' 
                else:
                    styles[i] = 'background-color: #ffcccc; color: black;' 
        return styles

    styled_df = df.style.apply(highlight_active_row, axis=1)
    edited_df = st.data_editor(styled_df, num_rows="dynamic", use_container_width=True, height=250)
    
    if st.button("Save to Excel", type="primary", key="save_blank"):
        try:
            # Safely load the workbook using openpyxl natively
            wb = openpyxl.load_workbook(EXCEL_FILE)
            
            # Find the original position of the sheet so we don't move it to the end
            if SHEET_NAME in wb.sheetnames:
                sheet_idx = wb.sheetnames.index(SHEET_NAME)
                del wb[SHEET_NAME] # Delete the old corrupted/outdated sheet
            else:
                sheet_idx = len(wb.sheetnames)
                
            # Create a fresh sheet in the exact same position
            ws = wb.create_sheet(SHEET_NAME, sheet_idx)
            
            # Write the updated pandas dataframe row by row cleanly
            for r in dataframe_to_rows(edited_df, index=False, header=True):
                ws.append(r)
                
            # Save the workbook
            wb.save(EXCEL_FILE)
            
            st.session_state.df = edited_df
            st.success("Successfully saved changes to Excel!")
            st.rerun()
        except Exception as e:
            st.error(f"Failed to save: {e}. Ensure the Excel file is closed on your computer.")

else:
    st.error(f"Could not load data. Ensure '{SHEET_NAME}' exists and has 'PRE_SNAP' and 'POST_SNAP' columns.")